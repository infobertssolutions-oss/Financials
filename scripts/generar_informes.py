#!/usr/bin/env python3
"""
Genera los estados financieros de Berts Solutions SL a partir de los CSV de /datos/.

Uso:
    python3 scripts/generar_informes.py                    # fecha de corte = hoy
    python3 scripts/generar_informes.py --corte 2026-06-30 # balance a una fecha concreta

Produce:
    informes/Estados_Financieros_Berts.xlsx  Balance, PyG, libro diario, IVA, vehículos... (con fórmulas)
    informes/balance_situacion.md            Lo mismo en texto, para leerlo en GitHub o en el móvil
    informes/cuenta_perdidas_ganancias.md
    informes/iva.md
    informes/rentabilidad_vehiculos.md
    informes/avisos.md                       Lo que falta o no cuadra

Requiere openpyxl (pip install openpyxl).
"""
import argparse
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

import contabilidad as cb

INFORMES = cb.BASE / "informes"
INFORMES.mkdir(exist_ok=True)
EXCEL = INFORMES / "Estados_Financieros_Berts.xlsx"

# --- estilo ---------------------------------------------------------------
FUENTE = "Arial"
F_NORMAL = Font(name=FUENTE, size=10)
F_NEGRITA = Font(name=FUENTE, size=10, bold=True)
F_TITULO = Font(name=FUENTE, size=14, bold=True, color="1F3864")
F_SUBTITULO = Font(name=FUENTE, size=9, italic=True, color="595959")
F_CABECERA = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
F_INPUT = Font(name=FUENTE, size=10, color="0000FF")
F_INPUT_NEGRITA = Font(name=FUENTE, size=10, color="0000FF", bold=True)
F_ENLACE = Font(name=FUENTE, size=10, color="008000")
F_ENLACE_NEGRITA = Font(name=FUENTE, size=10, color="008000", bold=True)
F_CODIGO = Font(name=FUENTE, size=8, color="808080")
RELLENO_CABECERA = PatternFill("solid", fgColor="1F3864")
RELLENO_TOTAL = PatternFill("solid", fgColor="D9E1F2")
RELLENO_SUBTOTAL = PatternFill("solid", fgColor="F2F2F2")
RELLENO_INPUT = PatternFill("solid", fgColor="FFFF00")
RELLENO_ALTA = PatternFill("solid", fgColor="F8CBAD")
RELLENO_MEDIA = PatternFill("solid", fgColor="FFE699")
RELLENO_BAJA = PatternFill("solid", fgColor="E2EFDA")
BORDE_TOTAL = Border(top=Side(style="thin"), bottom=Side(style="double"))
BORDE_SUB = Border(top=Side(style="thin"))
EUROS = '#,##0.00;(#,##0.00);"-"'
FECHA = "DD/MM/YYYY"

# Rangos del libro diario (se fijan en fijar_rangos() según el número de apuntes)
COL_FECHA = COL_CUENTA = COL_DEBE = COL_HABER = COL_TIPO = COL_VEH = COL_SALDO = COL_LBAL = COL_LPYG = ""


def fijar_rangos(ultima_fila):
    global COL_FECHA, COL_CUENTA, COL_DEBE, COL_HABER, COL_TIPO, COL_VEH, COL_SALDO, COL_LBAL, COL_LPYG
    r = lambda col: f"Diario!${col}$2:${col}${ultima_fila}"  # noqa: E731
    COL_FECHA, COL_CUENTA, COL_DEBE, COL_HABER = r("B"), r("C"), r("F"), r("G")
    COL_TIPO, COL_VEH, COL_SALDO, COL_LBAL, COL_LPYG = r("H"), r("I"), r("K"), r("L"), r("M")


eur = cb.eur


def cabecera_tabla(ws, fila, textos, anchos=None):
    for i, t in enumerate(textos, start=1):
        c = ws.cell(fila, i, t)
        c.font, c.fill = F_CABECERA, RELLENO_CABECERA
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if anchos:
        for i, a in enumerate(anchos, start=1):
            ws.column_dimensions[get_column_letter(i)].width = a


def titulo(ws, texto, subtitulo):
    ws["A1"] = texto
    ws["A1"].font = F_TITULO
    ws["A2"] = subtitulo
    ws["A2"].font = F_SUBTITULO


def estilo_linea(ws, fila, nivel, ncols, es_total=False):
    for col in range(1, ncols + 1):
        c = ws.cell(fila, col)
        if col == 1:
            c.font = F_CODIGO
        elif nivel == 0 or es_total:
            c.font = F_NEGRITA
        else:
            c.font = F_NORMAL
        if col >= 3:
            c.number_format = EUROS
        if es_total:
            c.fill, c.border = RELLENO_TOTAL, BORDE_TOTAL
        elif nivel == 0:
            c.fill, c.border = RELLENO_SUBTOTAL, BORDE_SUB
    ws.cell(fila, 2).alignment = Alignment(indent=nivel * 2)


# --- hojas ------------------------------------------------------------------
def hoja_plan(wb, cont):
    ws = wb.create_sheet("Plan de cuentas")
    titulo(ws, "Plan de cuentas utilizado (PGC Pymes)",
           "Relaciona cada cuenta con su línea del Balance y de la PyG. El Diario la usa para clasificar cada apunte.")
    cabecera_tabla(ws, 4, ["Cuenta", "Nombre", "Línea balance", "Línea PyG", "Texto línea balance",
                           "Texto línea PyG"], [10, 55, 13, 11, 55, 55])
    textos_bal = {l[0]: l[1] for l in cb.BALANCE_ACTIVO + cb.BALANCE_PASIVO}
    textos_bal.update({"D551": "Socios (deudor o acreedor según saldo)",
                       "D555": "Partidas pendientes de aplicación (deudor o acreedor según saldo)"})
    textos_pyg = {l[0]: l[1] for l in cb.PYG}
    cuentas = sorted(set(cb.PLAN_CUENTAS) | {a.cuenta for a in cont.apuntes})
    for i, cuenta in enumerate(cuentas, start=5):
        lb, lp = cb.linea_balance(cuenta), cb.linea_pyg(cuenta) or "-"
        valores = [cuenta, cb.nombre_cuenta(cuenta), lb, lp, textos_bal.get(lb, ""), textos_pyg.get(lp, "")]
        for j, v in enumerate(valores, start=1):
            c = ws.cell(i, j, v)
            c.font = F_NORMAL
            if j == 1:
                c.number_format = "@"
    ws.freeze_panes = "A5"
    return ws


def hoja_diario(wb, cont):
    ws = wb.create_sheet("Diario")
    cabecera_tabla(ws, 1, ["Asiento", "Fecha", "Cuenta", "Nombre de la cuenta", "Concepto", "Debe (€)",
                           "Haber (€)", "Tipo", "Vehículo", "Documento", "Saldo D−H (€)", "Línea balance",
                           "Línea PyG"], [8, 11, 8, 34, 60, 12, 12, 6, 8, 40, 13, 10, 9])
    n_plan = 4 + len(set(cb.PLAN_CUENTAS) | {a.cuenta for a in cont.apuntes})
    plan_a, plan_c, plan_d = (f"'Plan de cuentas'!${c}$5:${c}${n_plan}" for c in "ACD")
    for i, a in enumerate(cont.apuntes, start=2):
        valores = [a.asiento, a.fecha, a.cuenta, cb.nombre_cuenta(a.cuenta), a.concepto,
                   a.debe or None, a.haber or None, a.tipo, a.vehiculo, a.documento]
        for j, v in enumerate(valores, start=1):
            c = ws.cell(i, j, v)
            c.font = F_NORMAL
        ws.cell(i, 2).number_format = FECHA
        ws.cell(i, 3).number_format = "@"
        ws.cell(i, 6).number_format = EUROS
        ws.cell(i, 7).number_format = EUROS
        ws.cell(i, 11, f"=F{i}-G{i}").number_format = EUROS
        ws.cell(i, 12, f"=INDEX({plan_c},MATCH(C{i},{plan_a},0))")
        ws.cell(i, 13, f"=INDEX({plan_d},MATCH(C{i},{plan_a},0))")
        for j in (11, 12, 13):
            ws.cell(i, j).font = F_NORMAL
    ultima = len(cont.apuntes) + 1
    fila_total = ultima + 2
    ws.cell(fila_total, 5, "SUMAS").font = F_NEGRITA
    for col in ("F", "G", "K"):
        c = ws[f"{col}{fila_total}"]
        c.value = f"=SUM({col}2:{col}{ultima})"
        c.font, c.number_format, c.border = F_NEGRITA, EUROS, BORDE_TOTAL
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:M{ultima}"
    return fila_total


def sumifs_saldo(clave, columna_linea, condiciones_fecha, signo=1, extra=""):
    formula = f"SUMIFS({COL_SALDO},{columna_linea},\"{clave}\",{condiciones_fecha}{extra})"
    return f"-{formula}" if signo < 0 else formula


def escribir_estado(ws, lineas, fila_inicio, columnas_fecha, pyg=False, valores_fijos=None):
    """Escribe las líneas de un estado. columnas_fecha: {col: condición de fecha para SUMIFS}.
    valores_fijos: {col: {codigo: valor}} para columnas con datos externos (Fisama)."""
    filas = {}
    for n, (codigo, _, _, _) in enumerate(lineas):
        filas[codigo] = fila_inicio + n
    col_linea = COL_LPYG if pyg else COL_LBAL
    extra = f",{COL_TIPO},\"<>REG\"" if pyg else ""
    for codigo, texto, nivel, spec in lineas:
        f = filas[codigo]
        ws.cell(f, 1, codigo)
        ws.cell(f, 2, texto)
        es_total = codigo in ("TA", "TP", "A4")
        for col, cond in columnas_fecha.items():
            celda = ws.cell(f, col)
            fijos = (valores_fijos or {}).get(col)
            if spec[0] == "suma":
                celda.value = "=" + "+".join(f"{get_column_letter(col)}{filas[c]}" for c in spec[1])
            elif fijos is not None:
                celda.value = fijos.get(codigo, 0)
            elif spec[0] == "cuentas":
                celda.value = "=" + sumifs_saldo(codigo, col_linea, cond, spec[2], extra)
            elif spec[0] == "deudor":
                celda.value = f"=MAX(0,{sumifs_saldo(spec[1], col_linea, cond)})"
            elif spec[0] == "acreedor":
                celda.value = f"=MAX(0,-{sumifs_saldo(spec[1], col_linea, cond)})"
        estilo_linea(ws, f, nivel, 2 + len(columnas_fecha), es_total)
        for col in (valores_fijos or {}):
            if spec[0] != "suma":
                ws.cell(f, col).font = F_INPUT_NEGRITA if nivel == 0 else F_INPUT
    return filas


def hoja_balance(wb):
    ws = wb.create_sheet("Balance", 1)
    titulo(ws, "Balance de Situación - Berts Solutions SL (CIF B23928419)",
           "Modelo PGC Pymes. Todas las cifras salen del Diario con fórmulas. La columna del 31/12/2025 "
           "reproduce el balance presentado por Fisama.")
    cabecera_tabla(ws, 4, ["", "ACTIVO", "", ""], [7, 62, 16, 16])
    ws["C4"], ws["D4"] = "=FechaComparativa", "=FechaCorte"
    for c in ("C4", "D4"):
        ws[c].number_format = FECHA
    cond = {3: f"{COL_FECHA},\"<=\"&C$4", 4: f"{COL_FECHA},\"<=\"&D$4"}
    filas_a = escribir_estado(ws, cb.BALANCE_ACTIVO, 5, cond)
    fila_p = max(filas_a.values()) + 2
    cabecera_tabla(ws, fila_p, ["", "PATRIMONIO NETO Y PASIVO", "", ""])
    ws.cell(fila_p, 3, "=FechaComparativa").number_format = FECHA
    ws.cell(fila_p, 4, "=FechaCorte").number_format = FECHA
    filas_p = escribir_estado(ws, cb.BALANCE_PASIVO, fila_p + 1, cond)
    f = max(filas_p.values()) + 2
    ws.cell(f, 2, "Comprobación: Activo − (Patrimonio neto + Pasivo)").font = F_NEGRITA
    for col in ("C", "D"):
        c = ws[f"{col}{f}"]
        c.value = f"={col}{filas_a['TA']}-{col}{filas_p['TP']}"
        c.number_format, c.font = EUROS, F_NEGRITA
    ws.cell(f + 1, 2, "Si sale distinto de «-» (cero), algo está mal registrado.").font = F_SUBTITULO
    ws.freeze_panes = "C5"
    filas = dict(filas_a, **filas_p)
    filas["_check"] = f
    return filas


def hoja_pyg(wb, cont):
    ws = wb.create_sheet("PyG", 2)
    titulo(ws, "Cuenta de Pérdidas y Ganancias - Berts Solutions SL",
           "Modelo PGC Pymes. Ingresos en positivo, gastos entre paréntesis. Trimestres del año de la fecha de corte.")
    cabecera_tabla(ws, 4, ["", "", "2025 (Fisama)", "T1", "T2", "T3", "T4", "Acumulado a fecha de corte"],
                   [7, 62, 14, 14, 14, 14, 14, 16])
    for t, col in enumerate((4, 5, 6, 7), start=1):
        ws.cell(4, col, f"=\"T{t} \"&YEAR(FechaCorte)")
        ws.cell(5, col, f"=DATE(YEAR(FechaCorte),{3 * t - 2},1)")
        ws.cell(6, col, f"=MIN(DATE(YEAR(FechaCorte),{3 * t + 1},1)-1,FechaCorte)")
    ws.cell(5, 8, "=DATE(YEAR(FechaCorte),1,1)")
    ws.cell(6, 8, "=FechaCorte")
    ws.cell(5, 2, "Desde").font = F_SUBTITULO
    ws.cell(6, 2, "Hasta").font = F_SUBTITULO
    for col in range(4, 9):
        for fila in (5, 6):
            ws.cell(fila, col).number_format = FECHA
            ws.cell(fila, col).font = F_SUBTITULO
    ws.cell(5, 3, "01/01/2025").font = F_SUBTITULO
    ws.cell(6, 3, "31/12/2025").font = F_SUBTITULO
    cond = {col: f"{COL_FECHA},\">=\"&{get_column_letter(col)}$5,{COL_FECHA},\"<=\"&{get_column_letter(col)}$6"
            for col in range(3, 9)}
    filas = escribir_estado(ws, cb.PYG, 8, cond, pyg=True, valores_fijos={3: cb.PYG_2025_FISAMA})
    ws.cell(filas["P1"], 3).comment = Comment(
        "Fuente: Cuenta de PyG 2025 presentada por Fisama (documentos/cuentas_anuales/). "
        "Amortización 0 porque Fisama no amortizó.", "Claude")

    # Detalle por cuenta: para ver en qué se va el dinero
    f = max(filas.values()) + 3
    cabecera_tabla(ws, f, ["Cuenta", "Detalle por cuenta (ingresos y gastos)", "", "T1", "T2", "T3", "T4",
                           "Acumulado"])
    cuentas = sorted({a.cuenta for a in cont.apuntes if a.cuenta[:1] in ("6", "7")})
    fila_detalle_inicio = f + 1
    for i, cuenta in enumerate(cuentas, start=f + 1):
        ws.cell(i, 1, cuenta).number_format = "@"
        ws.cell(i, 2, cb.nombre_cuenta(cuenta))
        for col in range(4, 9):
            letra = get_column_letter(col)
            ws.cell(i, col, f"=-SUMIFS({COL_SALDO},{COL_CUENTA},$A{i},{COL_FECHA},\">=\"&{letra}$5,"
                            f"{COL_FECHA},\"<=\"&{letra}$6,{COL_TIPO},\"<>REG\")")
        estilo_linea(ws, i, 1, 8)
    fila_suma = fila_detalle_inicio + len(cuentas)
    ws.cell(fila_suma, 2, "Total = Resultado del ejercicio")
    for col in range(4, 9):
        letra = get_column_letter(col)
        ws.cell(fila_suma, col, f"=SUM({letra}{fila_detalle_inicio}:{letra}{fila_suma - 1})")
    estilo_linea(ws, fila_suma, 0, 8, es_total=True)
    ws.freeze_panes = "C7"
    return filas


def hoja_sumas_saldos(wb, cont):
    ws = wb.create_sheet("Sumas y saldos")
    titulo(ws, "Balance de sumas y saldos", "Movimientos acumulados de cada cuenta hasta la fecha de corte.")
    cabecera_tabla(ws, 4, ["Cuenta", "Nombre", "Debe (€)", "Haber (€)", "Saldo deudor (€)",
                           "Saldo acreedor (€)"], [9, 50, 14, 14, 15, 15])
    cuentas = sorted({a.cuenta for a in cont.apuntes})
    for i, cuenta in enumerate(cuentas, start=5):
        ws.cell(i, 1, cuenta).number_format = "@"
        ws.cell(i, 2, cb.nombre_cuenta(cuenta))
        ws.cell(i, 3, f"=SUMIFS({COL_DEBE},{COL_CUENTA},$A{i},{COL_FECHA},\"<=\"&FechaCorte)")
        ws.cell(i, 4, f"=SUMIFS({COL_HABER},{COL_CUENTA},$A{i},{COL_FECHA},\"<=\"&FechaCorte)")
        ws.cell(i, 5, f"=MAX(0,C{i}-D{i})")
        ws.cell(i, 6, f"=MAX(0,D{i}-C{i})")
        estilo_linea(ws, i, 1, 6)
    ult = 4 + len(cuentas)
    ws.cell(ult + 1, 2, "TOTALES")
    for col in "CDEF":
        ws[f"{col}{ult + 1}"] = f"=SUM({col}5:{col}{ult})"
    estilo_linea(ws, ult + 1, 0, 6, es_total=True)
    ws.freeze_panes = "C5"


def hoja_vehiculos(wb, cont):
    ws = wb.create_sheet("Vehículos")
    titulo(ws, "Inmovilizado: vehículos, amortización y rentabilidad",
           "Coste activado = precio + gastos necesarios para ponerlo a rodar (ITP, gestoría, aduana, matriculación). "
           "Amortización lineal por días.")
    cols = ["ID", "Vehículo", "Matrícula", "Estado", "Precio compra (€)", "¿Precio estimado?", "% amort. anual",
            "Coste activado total (€)", "Amortización dotada (€)", "Valor neto contable (€)", "Ingresos (€)",
            "Gastos incl. amortización (€)", "Resultado del vehículo (€)"]
    cabecera_tabla(ws, 4, cols, [5, 30, 18, 16, 13, 10, 9, 14, 14, 14, 13, 14, 14])
    ws.row_dimensions[4].height = 42
    hasta = f"{COL_FECHA},\"<=\"&FechaCorte"
    for i, v in enumerate(cont.vehiculos, start=5):
        vid = v["vehiculo_id"]
        datos = [vid, v["marca_modelo"], v["matricula"], v["estado_actual"], cb.num(v, "precio_compra_eur"),
                 v.get("precio_es_estimado_si_no", ""), cb.num(v, "pct_amortizacion_anual") / 100]
        for j, x in enumerate(datos, start=1):
            ws.cell(i, j, x).font = F_INPUT if j in (5, 6, 7) else F_NORMAL
        ws.cell(i, 7).number_format = "0%"
        if not v.get("fecha_inicio_amortizacion"):
            ws.cell(i, 7).comment = Comment("No se amortiza: todavía no está en condiciones de uso.", "Claude")
        f_veh = f"{COL_VEH},$A{i}"
        ws.cell(i, 8, f"=SUMIFS({COL_DEBE},{f_veh},{COL_CUENTA},\"218*\",{hasta})")
        ws.cell(i, 9, f"=SUMIFS({COL_SALDO},{f_veh},{COL_CUENTA},\"681*\",{hasta})")
        ws.cell(i, 10, f"=SUMIFS({COL_SALDO},{f_veh},{COL_CUENTA},\"218*\",{hasta})"
                       f"+SUMIFS({COL_SALDO},{f_veh},{COL_CUENTA},\"281*\",{hasta})")
        ws.cell(i, 11, f"=-SUMIFS({COL_SALDO},{f_veh},{COL_CUENTA},\"7*\",{COL_TIPO},\"<>REG\",{hasta})")
        ws.cell(i, 12, f"=SUMIFS({COL_SALDO},{f_veh},{COL_CUENTA},\"6*\",{COL_TIPO},\"<>REG\",{hasta})")
        ws.cell(i, 13, f"=K{i}-L{i}")
        for col in range(8, 14):
            ws.cell(i, col).font = F_NORMAL
        for col in (5, *range(8, 14)):
            ws.cell(i, col).number_format = EUROS
    ult = 4 + len(cont.vehiculos)
    ws.cell(ult + 1, 2, "TOTAL")
    for col in (5, *range(8, 14)):
        letra = get_column_letter(col)
        ws.cell(ult + 1, col, f"=SUM({letra}5:{letra}{ult})")
    estilo_linea(ws, ult + 1, 0, 13, es_total=True)
    for col in (1, 3, 4, 6, 7):
        ws.cell(ult + 1, col).number_format = "General"
    nota = ult + 3
    ws.cell(nota, 2, "Gastos sin vehículo asignado (hosting, combustible sin identificar...) no aparecen aquí.").font = \
        F_SUBTITULO
    ws.cell(nota + 1, 2, "El Ford Fiesta (V3) vendido: su «Gastos» incluye la pérdida contable de la venta.").font = \
        F_SUBTITULO
    ws.freeze_panes = "C5"


def hoja_iva(wb, cont):
    ws = wb.create_sheet("IVA")
    titulo(ws, "IVA por trimestre (para contrastar con el modelo 303 de Fisama)",
           "Estimación a partir de las facturas registradas. El IVA no deducible (facturas a nombre de otra persona, "
           "tickets) no cuenta como soportado.")
    cabecera_tabla(ws, 4, ["Trimestre", "Desde", "Hasta", "IVA repercutido (€)", "IVA soportado deducible (€)",
                           "Diferencia del trimestre (€)", "Estado"], [12, 12, 12, 17, 18, 18, 60])
    liquidados = {l["trimestre"]: l for l in cont.liquidaciones_iva}
    anio = cont.corte.year
    for t in range(1, 5):
        i = 4 + t
        etiqueta = f"{anio}-T{t}"
        ws.cell(i, 1, etiqueta)
        ws.cell(i, 2, date(anio, 3 * t - 2, 1)).number_format = FECHA
        ws.cell(i, 3, cb.fin_de_trimestre(anio, t)).number_format = FECHA
        rango = f"{COL_FECHA},\">=\"&B{i},{COL_FECHA},\"<=\"&C{i},{COL_TIPO},\"<>IVA\""
        ws.cell(i, 4, f"=-SUMIFS({COL_SALDO},{COL_CUENTA},\"477*\",{rango})")
        ws.cell(i, 5, f"=SUMIFS({COL_SALDO},{COL_CUENTA},\"472*\",{rango})")
        ws.cell(i, 6, f"=D{i}-E{i}")
        if etiqueta in liquidados:
            l = liquidados[etiqueta]
            if l["resultado"] >= 0:
                estado = (f"Cerrado. Tras compensar {eur(l['compensado'])} de trimestres anteriores: "
                          f"{eur(l['resultado'])} A INGRESAR")
            else:
                estado = f"Cerrado. {eur(-l['resultado'])} a compensar en trimestres siguientes"
        elif cb.fin_de_trimestre(anio, t) >= cont.corte >= date(anio, 3 * t - 2, 1):
            estado = "En curso"
        else:
            estado = "-"
        ws.cell(i, 7, estado)
        estilo_linea(ws, i, 1, 7)
        ws.cell(i, 2).number_format = ws.cell(i, 3).number_format = FECHA
        ws.cell(i, 7).number_format = "General"
    ws.cell(9, 1, "TOTAL")
    for col in "DEF":
        ws[f"{col}9"] = f"=SUM({col}5:{col}8)"
    estilo_linea(ws, 9, 0, 7, es_total=True)
    ws.cell(9, 2).number_format = ws.cell(9, 3).number_format = ws.cell(9, 7).number_format = "General"


def hoja_avisos(wb, cont):
    ws = wb.create_sheet("Pendiente")
    titulo(ws, "Qué falta o no cuadra", "Se genera solo cada vez que se actualizan los datos. "
                                         "Cuando nos paséis la información, el aviso desaparece.")
    cabecera_tabla(ws, 4, ["Prioridad", "Tema", "Detalle"], [10, 32, 120])
    rellenos = {"ALTA": RELLENO_ALTA, "MEDIA": RELLENO_MEDIA, "BAJA": RELLENO_BAJA}
    for i, a in enumerate(cont.avisos, start=5):
        for j, v in enumerate((a.prioridad, a.tema, a.detalle), start=1):
            c = ws.cell(i, j, v)
            c.font = F_NORMAL
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(i, 1).fill = rellenos[a.prioridad]
    ws.freeze_panes = "A5"


def hoja_resumen(wb, cont, filas_bal, filas_pyg, fila_total_diario):
    ws = wb.active
    ws.title = "Resumen"
    titulo(ws, "Estados financieros - Berts Solutions SL",
           f"Generado el {date.today():%d/%m/%Y} con scripts/generar_informes.py a partir de /datos/. "
           "No se edita a mano: se vuelve a generar con cada actualización.")
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 52
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 70

    ws["B4"], ws["C4"] = "Fecha de corte del balance", cont.corte
    ws["B5"], ws["C5"] = "Fecha comparativa (último cierre)", cb.FECHA_APERTURA
    for c in ("C4", "C5"):
        ws[c].number_format = FECHA
        ws[c].font = F_INPUT_NEGRITA
        ws[c].fill = RELLENO_INPUT
    ws["D4"] = ("Puedes poner una fecha ANTERIOR para ver el balance a ese día. Para una fecha posterior, "
                "vuelve a generar el informe (la amortización se calcula hasta la fecha de generación).")
    ws["D4"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["D4"].font = F_SUBTITULO
    ws.row_dimensions[4].height = 36
    wb.defined_names["FechaCorte"] = DefinedName("FechaCorte", attr_text="Resumen!$C$4")
    wb.defined_names["FechaComparativa"] = DefinedName("FechaComparativa", attr_text="Resumen!$C$5")

    f = 7
    cabecera_tabla(ws, f, ["", "Cifras clave a la fecha de corte", "Importe (€)", "Qué significa"])
    clave = [
        ("Total activo (lo que tiene la empresa)", f"=Balance!D{filas_bal['TA']}",
         "Vehículos (neto de amortización) + dinero + lo que le deben."),
        ("Tesorería según contabilidad", f"=Balance!D{filas_bal['BVII']}",
         "Debería coincidir con el extracto del BBVA. Sin extractos no está comprobado."),
        ("Vehículos (valor neto contable)", f"=Balance!D{filas_bal['AII']}",
         "Coste de los vehículos menos la amortización acumulada."),
        ("Clientes pendientes de cobro", f"=Balance!D{filas_bal['BIII1']}",
         "Facturas emitidas que no nos consta que se hayan cobrado."),
        ("Patrimonio neto (lo que es de los socios)", f"=Balance!D{filas_bal['PA']}",
         "Capital + aportaciones + resultados acumulados."),
        ("Partidas pendientes de aplicación (cuenta 555)", f"=Balance!D{filas_bal['PCIII3']}",
         "Pagos cuyo origen desconocemos. Tiene que llegar a 0 cuando tengamos los extractos."),
        ("Ingresos por alquiler del ejercicio", f"=PyG!H{filas_pyg['P1']}", "Cifra de negocios acumulada."),
        ("Resultado del ejercicio (beneficio / pérdida)", f"=PyG!H{filas_pyg['A4']}",
         "Ingresos menos todos los gastos, incluida la amortización y la pérdida del Ford Fiesta."),
    ]
    for i, (texto, formula, explicacion) in enumerate(clave, start=f + 1):
        ws.cell(i, 2, texto).font = F_NORMAL
        c = ws.cell(i, 3, formula)
        c.font, c.number_format = F_ENLACE, EUROS
        ws.cell(i, 4, explicacion).font = F_SUBTITULO

    f = f + len(clave) + 2
    cabecera_tabla(ws, f, ["", "Controles de cuadre", "Diferencia (€)", "Estado"])
    controles = [
        ("Activo = Patrimonio neto + Pasivo", f"=Balance!D{filas_bal['TA']}-Balance!D{filas_bal['TP']}"),
        ("Libro diario: Debe = Haber", f"=Diario!F{fila_total_diario}-Diario!G{fila_total_diario}"),
        ("Resultado de la PyG = Resultado del balance", f"=PyG!H{filas_pyg['A4']}-Balance!D{filas_bal['PAVII']}"),
        ("Balance 31/12/2025 = el de Fisama (75.318,61 €)", f"=Balance!C{filas_bal['TA']}-75318.61"),
    ]
    for i, (texto, formula) in enumerate(controles, start=f + 1):
        ws.cell(i, 2, texto).font = F_NORMAL
        c = ws.cell(i, 3, formula)
        c.font, c.number_format = F_ENLACE, EUROS
        ws.cell(i, 4, f"=IF(ABS(C{i})<0.01,\"OK\",\"ERROR - revisar\")").font = F_NEGRITA

    f = f + len(controles) + 2
    cabecera_tabla(ws, f, ["", "Hojas de este libro", "", ""])
    hojas = [
        ("Balance", "Balance de situación (modelo PGC Pymes) a 31/12/2025 y a la fecha de corte."),
        ("PyG", "Cuenta de pérdidas y ganancias por trimestres + detalle por cuenta."),
        ("Vehículos", "Coste, amortización, valor neto y resultado de cada vehículo."),
        ("IVA", "IVA repercutido y soportado por trimestre, para contrastar con el modelo 303."),
        ("Pendiente", f"{len(cont.avisos)} avisos: información que falta o no cuadra."),
        ("Diario", "Libro diario: todos los asientos. Todo lo demás sale de aquí."),
        ("Sumas y saldos", "Saldo de cada cuenta contable."),
        ("Plan de cuentas", "Qué cuenta va a qué línea del balance y de la PyG."),
    ]
    for i, (hoja, texto) in enumerate(hojas, start=f + 1):
        ws.cell(i, 2, hoja).font = F_NEGRITA
        ws.cell(i, 4, texto).font = F_NORMAL

    f = f + len(hojas) + 2
    ws.cell(f, 2, "Leyenda de colores").font = F_NEGRITA
    leyenda = [
        (F_INPUT, None, "Azul: dato introducido (fecha de corte, cifras de Fisama, precios de compra)."),
        (F_NORMAL, None, "Negro: fórmula calculada a partir del Diario."),
        (F_ENLACE, None, "Verde: enlace a otra hoja de este libro."),
        (F_INPUT, RELLENO_INPUT, "Fondo amarillo: celda que puedes cambiar."),
    ]
    for i, (fuente, relleno, texto) in enumerate(leyenda, start=f + 1):
        c = ws.cell(i, 2, texto)
        c.font = fuente
        if relleno:
            c.fill = relleno


def generar_excel(cont):
    wb = Workbook()
    fijar_rangos(len(cont.apuntes) + 1)
    hoja_plan(wb, cont)
    fila_total_diario = hoja_diario(wb, cont)
    filas_bal = hoja_balance(wb)
    filas_pyg = hoja_pyg(wb, cont)
    hoja_vehiculos(wb, cont)
    hoja_iva(wb, cont)
    hoja_avisos(wb, cont)
    hoja_sumas_saldos(wb, cont)
    hoja_resumen(wb, cont, filas_bal, filas_pyg, fila_total_diario)
    orden = ["Resumen", "Balance", "PyG", "Vehículos", "IVA", "Pendiente", "Diario", "Sumas y saldos",
             "Plan de cuentas"]
    wb._sheets = [wb[n] for n in orden]
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = ws.title in ("Diario", "Sumas y saldos", "Plan de cuentas")
    wb.save(EXCEL)


# --- informes en texto --------------------------------------------------------
def cabecera_md(titulo_md, cont):
    return (f"# {titulo_md}\n\n_Generado automáticamente el {date.today():%d/%m/%Y} con datos hasta el "
            f"{cont.corte:%d/%m/%Y}. No editar a mano: se sobrescribe. La versión completa, con fórmulas, está en "
            f"[`Estados_Financieros_Berts.xlsx`](Estados_Financieros_Berts.xlsx)._\n\n")


def tabla_estado(lineas, columnas):
    """columnas: [(titulo, {codigo: valor})]"""
    out = "| | " + " | ".join(t for t, _ in columnas) + " |\n"
    out += "|---|" + "---:|" * len(columnas) + "\n"
    for codigo, texto, nivel, _ in lineas:
        sangria = "&nbsp;&nbsp;&nbsp;&nbsp;" * nivel
        negrita = nivel == 0
        celdas = [eur(v[codigo]) if v.get(codigo) is not None else "-" for _, v in columnas]
        if negrita:
            texto = f"**{texto}**"
            celdas = [f"**{c}**" for c in celdas]
        out += f"| {sangria}{texto} | " + " | ".join(celdas) + " |\n"
    return out


def md_balance(cont):
    bal_ini = cont.valor_lineas(cb.BALANCE_ACTIVO + cb.BALANCE_PASIVO, cb.FECHA_APERTURA)
    bal = cont.valor_lineas(cb.BALANCE_ACTIVO + cb.BALANCE_PASIVO, cont.corte)
    cols = [(f"{cb.FECHA_APERTURA:%d/%m/%Y}", bal_ini), (f"{cont.corte:%d/%m/%Y}", bal)]
    out = cabecera_md("Balance de Situación", cont)
    out += "## Activo\n\n" + tabla_estado(cb.BALANCE_ACTIVO, cols)
    out += "\n## Patrimonio neto y pasivo\n\n" + tabla_estado(cb.BALANCE_PASIVO, cols)
    dif = round(bal["TA"] - bal["TP"], 2)
    out += f"\n**Comprobación:** Activo − (PN + Pasivo) = {eur(dif)} {'✅' if not dif else '❌'}\n"
    pend = bal["PCIII3"]
    if pend:
        out += (f"\n> ⚠️ **{eur(pend)} en «partidas pendientes de aplicación»**: compras (sobre todo de vehículos) "
                "cuyo pago no sabemos de dónde salió. Si se pagaron desde el BBVA, la tesorería real es menor que la "
                "que aparece; si las pagasteis vosotros, son deuda de la empresa con vosotros o una aportación. "
                "Se resuelve con los extractos bancarios.\n")
    (INFORMES / "balance_situacion.md").write_text(out, encoding="utf-8")


def md_pyg(cont):
    anio = cont.corte.year
    cols = [("2025 (Fisama)", {**{l[0]: 0.0 for l in cb.PYG}, **cb.PYG_2025_FISAMA})]
    pyg_2025 = cols[0][1]
    for t in range(1, 5):
        desde = date(anio, 3 * t - 2, 1)
        if desde > cont.corte:
            break
        hasta = min(cb.fin_de_trimestre(anio, t), cont.corte)
        cols.append((f"T{t} {anio}", cont.valor_lineas(cb.PYG, hasta, desde, pyg=True)))
    cols.append((f"Acumulado {anio}", cont.valor_lineas(cb.PYG, cont.corte, date(anio, 1, 1), pyg=True)))
    # Subtotales de la columna de Fisama
    subt = {"P7": pyg_2025["P7a"], "A1": pyg_2025["P1"] + pyg_2025["P7a"]}
    subt.update({"A3": subt["A1"], "A4": subt["A1"]})
    pyg_2025.update(subt)
    out = cabecera_md("Cuenta de Pérdidas y Ganancias", cont)
    out += tabla_estado(cb.PYG, cols)
    out += "\n## Detalle por cuenta (acumulado del ejercicio)\n\n| Cuenta | Concepto | Importe |\n|---|---|---:|\n"
    saldos = {}
    for a in cont.apuntes:
        if a.cuenta[:1] in ("6", "7") and a.tipo != "REG" and date(anio, 1, 1) <= a.fecha <= cont.corte:
            saldos[a.cuenta] = saldos.get(a.cuenta, 0) - a.saldo
    for cuenta in sorted(saldos):
        out += f"| {cuenta} | {cb.nombre_cuenta(cuenta)} | {eur(saldos[cuenta])} |\n"
    for v in cont.ventas_inmovilizado:
        out += (f"\n> La venta de **{v['vehiculo']}** el {v['fecha']:%d/%m/%Y} por {eur(v['precio'])} (+IVA), con un "
                f"coste de {eur(v['coste'])} y {eur(v['amortizacion'])} ya amortizados, da un resultado de "
                f"**{eur(v['resultado'])}** (línea 11 de la PyG). No es un ingreso ordinario.\n")
    (INFORMES / "cuenta_perdidas_ganancias.md").write_text(out, encoding="utf-8")


def md_iva(cont):
    out = cabecera_md("IVA por trimestre", cont)
    out += ("Estimación para contrastar con el modelo 303 que presente la gestoría. El IVA no deducible "
            "(facturas a nombre de otra persona, tickets sin datos fiscales) no se cuenta como soportado.\n\n")
    out += "| Trimestre | Repercutido | Soportado deducible | Compensado de trimestres anteriores | Resultado 303 |\n"
    out += "|---|---:|---:|---:|---|\n"
    for l in cont.liquidaciones_iva:
        res = (f"**{eur(l['resultado'])} a ingresar**" if l["resultado"] >= 0
               else f"{eur(-l['resultado'])} a compensar")
        out += (f"| {l['trimestre']} | {eur(l['repercutido'])} | {eur(l['soportado'])} | {eur(l['compensado'])} "
                f"| {res} |\n")
    rep = -cont.saldo(cb.CUENTA_IVA_REPERCUTIDO, hasta=cont.corte)
    sop = cont.saldo(cb.CUENTA_IVA_SOPORTADO, hasta=cont.corte)
    comp = cont.saldo(cb.CUENTA_IVA_A_COMPENSAR, hasta=cont.corte)
    out += (f"| {cb.trimestre(cont.corte)} (en curso) | {eur(rep)} | {eur(sop)} | {eur(comp)} "
            f"| {eur(rep - sop - comp)} provisional |\n")
    (INFORMES / "iva.md").write_text(out, encoding="utf-8")


def md_vehiculos(cont):
    out = cabecera_md("Vehículos: coste, amortización y rentabilidad", cont)
    out += ("| Vehículo | Estado | Coste activado | Amortizado | Valor neto | Ingresos | Gastos (incl. amortización) "
            "| Resultado |\n|---|---|---:|---:|---:|---:|---:|---:|\n")
    for v in cont.vehiculos:
        vid = v["vehiculo_id"]
        coste = sum(a.debe for a in cont.apuntes if a.vehiculo == vid and a.cuenta.startswith("218")
                    and a.fecha <= cont.corte)
        amort = cont.saldo("681", hasta=cont.corte, vehiculo=vid)
        vnc = cont.saldo(["218", "281"], hasta=cont.corte, vehiculo=vid)
        ing = -cont.saldo("7", hasta=cont.corte, vehiculo=vid, excluir_tipos=("REG",))
        gas = cont.saldo("6", hasta=cont.corte, vehiculo=vid, excluir_tipos=("REG",))
        estimado = " (precio estimado)" if v.get("precio_es_estimado_si_no") == "si" else ""
        out += (f"| {vid} {v['marca_modelo']} {v['matricula']}{estimado} | {v['estado_actual']} | {eur(coste)} "
                f"| {eur(amort)} | {eur(vnc)} | {eur(ing)} | {eur(gas)} | **{eur(ing - gas)}** |\n")
    out += ("\nEl coste activado incluye el precio y los gastos para poner el vehículo en marcha (ITP, gestoría, "
            "aduana, arancel, transporte, matriculación). Esos gastos **no** son gasto del año: se reparten en la vida "
            "útil a través de la amortización.\n")
    (INFORMES / "rentabilidad_vehiculos.md").write_text(out, encoding="utf-8")


def md_avisos(cont):
    out = cabecera_md("Qué falta o no cuadra", cont)
    iconos = {"ALTA": "🔴", "MEDIA": "🟠", "BAJA": "🟢"}
    tema_actual = None
    for a in cont.avisos:
        if a.tema != tema_actual:
            out += f"\n### {iconos[a.prioridad]} {a.tema}\n\n"
            tema_actual = a.tema
        out += f"- {a.detalle}\n"
    (INFORMES / "avisos.md").write_text(out, encoding="utf-8")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--corte", help="Fecha de corte AAAA-MM-DD (por defecto, hoy)")
    args = parser.parse_args()
    corte = date.fromisoformat(args.corte) if args.corte else date.today()

    cont = cb.construir(corte)
    generar_excel(cont)
    md_balance(cont)
    md_pyg(cont)
    md_iva(cont)
    md_vehiculos(cont)
    md_avisos(cont)

    bal = cont.valor_lineas(cb.BALANCE_ACTIVO + cb.BALANCE_PASIVO, corte)
    pyg = cont.valor_lineas(cb.PYG, corte, date(corte.year, 1, 1), pyg=True)
    print(f"Estados financieros a {corte:%d/%m/%Y} generados en {INFORMES}")
    print(f"  Total activo:          {eur(bal['TA'])}")
    print(f"  Total PN + pasivo:     {eur(bal['TP'])}")
    print(f"  Resultado del ejercicio: {eur(pyg['A4'])}")
    print(f"  Asientos: {max((a.asiento for a in cont.apuntes), default=0)} · Avisos: {len(cont.avisos)}")


if __name__ == "__main__":
    main()
